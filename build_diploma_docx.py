"""
Генератор итоговой выпускной квалификационной работы (ВКР).

Тема: «Разработка бизнес-плана по производству товара
(на примере ООО «КУПИШУЗ», г. Москва)».

Документ собирается с учётом требований мастер-правил MainRule
(Times New Roman 14, межстрочный интервал 1,5, абзацный отступ 1,25 см,
поля 30/10/20/20 мм, чёрный цвет текста, заголовки структурных элементов
прописными и по центру) и в стиле «ВКР Гарден» (по-главное нумерование
таблиц, рисунков и формул; академический формальный язык; ссылки в
квадратных скобках). Анг­лоязычные сокращения заменены русскими.

Запуск:  python3 build_diploma_docx.py
Выход:   ВКР_Купишуз_бизнес_план.docx
"""

from __future__ import annotations

from copy import deepcopy

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


# ---------------------------------------------------------------------------
#                            ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ
# ---------------------------------------------------------------------------


def set_cell_borders(cell) -> None:
    """Задаёт тонкие чёрные границы (0,5 пт) у одной ячейки таблицы."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")  # 0,5 пт
        b.set(qn("w:color"), "000000")
        tc_borders.append(b)
    tc_pr.append(tc_borders)


def style_run(run, *, size: int = 14, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:cs"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_paragraph(
    doc,
    text: str = "",
    *,
    style: str | None = None,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_line_indent: float | None = 1.25,
    size: int = 14,
    bold: bool = False,
    italic: bool = False,
    line_spacing: float = 1.5,
    space_before: float = 0.0,
    space_after: float = 0.0,
    page_break_before: bool = False,
    keep_with_next: bool = False,
):
    if style is not None:
        p = doc.add_paragraph(style=style)
    else:
        p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.page_break_before = page_break_before
    pf.keep_with_next = keep_with_next
    if text:
        r = p.add_run(text)
        style_run(r, size=size, bold=bold, italic=italic)
    return p


def add_heading1(doc, text: str, *, page_break: bool = True):
    """Заголовок структурного элемента или главы (Заголовок 1)."""
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.page_break_before = page_break
    pf.keep_with_next = True
    r = p.add_run(text.upper())
    style_run(r, size=14, bold=True)
    add_paragraph(doc, "", first_line_indent=0)  # пустая строка после заголовка
    return p


def add_heading2(doc, text: str):
    """Заголовок параграфа (Заголовок 2)."""
    p = doc.add_paragraph(style="Heading 2")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.first_line_indent = Cm(1.25)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.page_break_before = False
    pf.keep_with_next = True
    r = p.add_run(text)
    style_run(r, size=14, bold=True)
    return p


def add_table_title(doc, text: str):
    """Заголовок таблицы (слева, кегль 14, без отступа)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(6)
    pf.space_after = Pt(0)
    pf.keep_with_next = True
    r = p.add_run(text)
    style_run(r, size=14)
    return p


def add_figure_caption(doc, text: str):
    """Подпись рисунка (по центру, 12 пт)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    r = p.add_run(text)
    style_run(r, size=12)
    return p


def add_table(doc, headers: list[str], rows: list[list[str]], *, header_bold: bool = True):
    """Создаёт таблицу с тонкими чёрными границами, кегль 12, межстрочный 1,0."""
    n_cols = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for j, hdr in enumerate(headers):
        cell = table.cell(0, j)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(hdr)
        style_run(r, size=12, bold=header_bold)
        set_cell_borders(cell)

    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val))
            style_run(r, size=12)
            set_cell_borders(cell)
    return table


def add_table_source(doc, text: str):
    """Источник под таблицей (12 пт, по ширине, без отступа)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    r = p.add_run(text)
    style_run(r, size=12, italic=True)
    return p


def add_formula(doc, body: str, number: str):
    """Формула: тело по центру, номер у правого поля."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    # Используем табуляцию: тело по центру, номер выровнен по правому краю
    tab_stops = pf.tab_stops
    tab_stops.add_tab_stop(Cm(16.0), alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    r1 = p.add_run(body)
    style_run(r1, size=14, italic=True)
    r2 = p.add_run("\t" + number)
    style_run(r2, size=14)
    return p


def configure_page(section) -> None:
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(30)
    section.right_margin = Mm(10)
    section.gutter = Mm(0)
    section.header_distance = Mm(12)
    section.footer_distance = Mm(12)


def add_page_number_footer(section, *, suppress: bool = False) -> None:
    """Добавляет номер страницы в нижний колонтитул по центру (без точки)."""
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = ""
    if suppress:
        return
    r = p.add_run()
    style_run(r, size=14)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE   \\* MERGEFORMAT"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r._element.append(fld_begin)
    r._element.append(instr)
    r._element.append(fld_end)


def add_toc(doc) -> None:
    """Вставляет авто-собираемое поле оглавления (обновляется в Word нажатием F9)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run()
    style_run(r, size=14)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Оглавление обновляется в Microsoft Word: правый клик → «Обновить поле → Обновить целиком»."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r._element.append(fld_begin)
    r._element.append(instr)
    r._element.append(fld_separate)
    r._element.append(placeholder)
    r._element.append(fld_end)


def setup_styles(doc) -> None:
    """Базовые настройки стилей под требования МФЮА."""
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.25)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    for name, size, bold in (
        ("Heading 1", 14, True),
        ("Heading 2", 14, True),
        ("Heading 3", 14, True),
    ):
        st = styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.keep_with_next = True

    # Гиперссылка — чёрная, без подчёркивания
    if "Hyperlink" in [s.name for s in styles]:
        try:
            hl = styles["Hyperlink"]
            hl.font.color.rgb = RGBColor(0, 0, 0)
            hl.font.underline = False
        except Exception:
            pass


# ---------------------------------------------------------------------------
#                                  ТИТУЛ
# ---------------------------------------------------------------------------


def build_title_page(doc) -> None:
    # Верхняя шапка
    for txt in (
        "Министерство науки и высшего образования Российской Федерации",
        "Аккредитованное образовательное частное учреждение",
        "высшего образования «Московский финансово-юридический университет МФЮА»",
        "Кафедра менеджмента",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(txt)
        style_run(r, size=14, bold=True)

    for _ in range(6):
        add_paragraph(doc, "", first_line_indent=0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА")
    style_run(r, size=16, bold=True)

    add_paragraph(doc, "", first_line_indent=0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(
        "на тему: «Разработка бизнес-плана по производству товара "
        "(на примере ООО «КУПИШУЗ», г. Москва)»"
    )
    style_run(r, size=14, bold=True)

    for _ in range(2):
        add_paragraph(doc, "", first_line_indent=0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("Специальность: 38.02.04 Коммерция (по отраслям)")
    style_run(r, size=14)

    for _ in range(4):
        add_paragraph(doc, "", first_line_indent=0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("Исполнитель: обучающийся выпускного курса")
    style_run(r, size=14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("__________________________________________ ___________")
    style_run(r, size=14)

    add_paragraph(doc, "", first_line_indent=0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("Научный руководитель: ___________________ ___________")
    style_run(r, size=14)

    for _ in range(6):
        add_paragraph(doc, "", first_line_indent=0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("Москва, 2026")
    style_run(r, size=14)


# ---------------------------------------------------------------------------
#                                  СОДЕРЖАНИЕ
# ---------------------------------------------------------------------------


def build_toc(doc) -> None:
    add_heading1(doc, "Оглавление")
    add_toc(doc)


# ---------------------------------------------------------------------------
#                                  ВВЕДЕНИЕ
# ---------------------------------------------------------------------------


INTRO_PARAGRAPHS = [
    "Современная экономика Российской Федерации переживает этап глубокой структурной перестройки, в которой импортозамещение, развитие отечественной производственной базы и формирование собственных торговых марок крупных розничных операторов превращаются из самостоятельного направления в системное условие сохранения конкурентоспособности. Уход с российского рынка ряда иностранных торговых марок в 2022–2024 годах, повышенная ключевая ставка Банка России и заметные колебания валютного курса обострили проблему долгосрочной устойчивости бизнес-моделей, основанных преимущественно на ввозе готовой продукции из-за рубежа. В этих условиях разработка экономически обоснованного бизнес-плана по запуску собственного производства товара выступает действенным инструментом снижения зависимости от внешних поставщиков, повышения маржинальности и формирования дополнительной потребительской ценности.",

    "Актуальность темы исследования обусловлена тем, что значительная часть крупных розничных компаний, работающих в сегменте одежды и обуви, исторически выстраивала логистику и закупки на основе передачи производственных функций сторонним организациям. Изменение макроэкономической конъюнктуры, переориентация поставщиков на азиатские рынки, удорожание трансграничных платежей и рост таможенных издержек обусловливают целесообразность частичного переноса производственных операций в Российскую Федерацию или в государства — участники Евразийского экономического союза. Одновременно развитие сегмента электронной торговли расширяет возможности по сбыту собственных торговых марок без привлечения сторонних оптовых посредников. Совокупность указанных факторов формирует устойчивый спрос на проектно обоснованные бизнес-планы запуска собственных производственных линий в составе крупных торговых организаций.",

    "Степень научной разработанности темы достаточно высока. Теоретические основы бизнес-планирования заложены в трудах И. Ансоффа, Г. Минцберга, М. Портера, П. Друкера, Ф. Котлера, а также отечественных учёных В. М. Попова, С. И. Ляпунова, В. П. Грузинова, В. А. Горемыкина, А. Б. Идрисова, Г. Б. Клейнера, В. В. Ковалева, Е. С. Стояновой, А. Д. Шеремета. Прикладные аспекты разработки бизнес-планов производственных проектов исследованы в работах И. И. Мазура, В. Д. Шапиро, Б. А. Колтынюка. Несмотря на обширную теоретическую базу, проблема комплексного бизнес-планирования при запуске собственного производства внутри сложившейся торговой структуры остаётся изученной фрагментарно, что подтверждает целесообразность настоящего исследования.",

    "Объектом исследования выступает финансово-хозяйственная деятельность общества с ограниченной ответственностью «КУПИШУЗ» (идентификационный номер налогоплательщика 7705935687, город Москва). Предметом исследования является совокупность теоретических, методических и практических положений, связанных с разработкой бизнес-плана по производству товара в условиях действующего торгового предприятия.",

    "Целью выпускной квалификационной работы является разработка экономически обоснованного бизнес-плана по производству собственной торговой марки обуви для ООО «КУПИШУЗ» и оценка целесообразности его реализации на основе общепринятых показателей эффективности инвестиционных проектов.",

    "Для достижения поставленной цели в работе решаются следующие задачи: раскрыть теоретическую сущность, функции, виды и принципы бизнес-планирования на предприятии; провести сравнительный анализ ведущих методических стандартов разработки бизнес-плана (методика Организации Объединённых Наций по промышленному развитию, методика Европейского банка реконструкции и развития, методики крупных международных консалтинговых компаний, Методические рекомендации Министерства экономики Российской Федерации); выявить отраслевые особенности бизнес-планирования при запуске собственного производства в составе крупного розничного оператора сегмента модной одежды и обуви; дать организационно-экономическую характеристику ООО «КУПИШУЗ» и проанализировать его финансовые результаты за 2023–2025 годы; провести стратегический анализ внешней и внутренней среды предприятия; разработать бизнес-план по производству собственной торговой марки обуви, включающий маркетинговый, производственный, организационный, финансовый разделы и оценку рисков; рассчитать показатели коммерческой эффективности проекта (чистый дисконтированный доход, внутреннюю норму доходности, индекс рентабельности инвестиций, дисконтированный срок окупаемости) и сформулировать предложения по их улучшению.",

    "В качестве рабочей гипотезы исследования принимается положение о том, что запуск ограниченной по объёму производственной линии собственной торговой марки обуви в составе действующей торговой экосистемы ООО «КУПИШУЗ» обеспечивает прирост маржинальной прибыли по выделяемому ассортименту не менее чем на восемь процентных пунктов относительно базового уровня закупочной маржи и формирует положительный чистый дисконтированный доход на пятилетнем горизонте планирования при ставке дисконтирования, соответствующей средневзвешенной стоимости капитала предприятия. Указанная гипотеза проверяется во второй и третьей главах работы на фактическом материале базового предприятия.",

    "Методологическую основу исследования составляют общенаучные методы (анализ, синтез, индукция, дедукция, аналогия), а также специальные методы экономических исследований: горизонтальный и вертикальный анализ бухгалтерской отчётности, коэффициентный анализ, стратегический анализ внутренней и внешней среды, методы дисконтирования денежных потоков, сценарный анализ рисков и метод оценки чувствительности.",

    "Информационную базу исследования составляют нормативные правовые акты Российской Федерации; данные Федеральной службы государственной статистики; отраслевые обзоры рынка обуви, опубликованные отечественными аналитическими агентствами; публичная бухгалтерская (финансовая) отчётность ООО «КУПИШУЗ» по формам Общероссийского классификатора управленческой документации 0710001 и 0710002 за 2023–2025 годы; сведения Единого государственного реестра юридических лиц, размещённые на портале Федеральной налоговой службы; материалы периодической печати и научных изданий.",

    "Теоретическая значимость работы заключается в систематизации подходов к бизнес-планированию производственных проектов, реализуемых в составе действующих торговых структур. Практическая значимость состоит в том, что разработанный бизнес-план может быть использован руководством ООО «КУПИШУЗ» при принятии управленческого решения о запуске собственного производства, а также адаптирован для применения в иных розничных компаниях со сходной бизнес-моделью.",

    "Структура работы. Выпускная квалификационная работа состоит из введения, трёх глав, заключения, списка использованных источников и приложений. Общий объём работы составляет 82 страницы машинописного текста, содержит 18 таблиц, 9 рисунков, 4 приложения. Список использованных источников включает 52 наименования.",
]


def build_intro(doc) -> None:
    add_heading1(doc, "Введение")
    for par in INTRO_PARAGRAPHS:
        add_paragraph(doc, par)


# ---------------------------------------------------------------------------
#                                    MAIN
# ---------------------------------------------------------------------------


def main() -> None:
    from content_ch1 import build_chapter1
    from content_ch2 import build_chapter2
    from content_ch3 import build_chapter3
    from content_back import build_conclusion, build_sources, build_appendices

    doc = Document()
    setup_styles(doc)

    section = doc.sections[0]
    configure_page(section)
    # На титульном листе номер не отображается
    section.different_first_page_header_footer = True
    add_page_number_footer(section)
    # Подавить номер на первой странице
    first_footer = section.first_page_footer
    fp = first_footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.text = ""

    build_title_page(doc)
    build_toc(doc)
    build_intro(doc)
    build_chapter1(doc)
    build_chapter2(doc)
    build_chapter3(doc)
    build_conclusion(doc)
    build_sources(doc)
    build_appendices(doc)

    output = "ВКР_Купишуз_бизнес_план.docx"
    doc.save(output)
    print(f"Документ сохранён: {output}")


if __name__ == "__main__":
    main()
